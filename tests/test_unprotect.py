"""Unprotect utility (SPEC-0002 AC-15..19)."""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
import pytest

from tests.builders import simple_workbook
from upl_helper.errors import UnsupportedWorkbookError
from upl_helper.profile.loader import file_sha256
from upl_helper.unprotect import unprotect_workbook

FAKE_VBA = b"\x00fake vba project bytes\xff" * 16


def _inject_extra_part(src: Path, dst: Path, name: str, payload: bytes) -> Path:
    """Copy a workbook, adding one extra zip entry.

    Used to stand in for parts openpyxl cannot create, such as a VBA project.
    """
    with (
        zipfile.ZipFile(src) as source,
        zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as out,
    ):
        for info in source.infolist():
            out.writestr(info, source.read(info.filename))
        out.writestr(name, payload)
    return dst


def test_protection_is_gone_from_the_output(tmp_path: Path) -> None:
    """AC-15."""
    book = simple_workbook(tmp_path / "p.xlsx", protect=True, protect_password="pw")
    report = unprotect_workbook(book, tmp_path / "out.xlsx")
    workbook = openpyxl.load_workbook(tmp_path / "out.xlsx")
    assert all(not sheet.protection.sheet for sheet in workbook.worksheets)
    assert report.removed_sheet_protection == ["Demonstration"]


def test_zip_entry_names_are_preserved(tmp_path: Path) -> None:
    """AC-16."""
    book = simple_workbook(tmp_path / "e.xlsx", protect=True, hidden_sheet=True)
    unprotect_workbook(book, tmp_path / "out.xlsx")
    with (
        zipfile.ZipFile(book) as before,
        zipfile.ZipFile(tmp_path / "out.xlsx") as after,
    ):
        assert set(before.namelist()) == set(after.namelist())


def test_vba_part_survives_byte_for_byte(tmp_path: Path) -> None:
    """AC-17."""
    base = simple_workbook(tmp_path / "base.xlsx", protect=True)
    macro_book = _inject_extra_part(
        base, tmp_path / "macro.xlsm", "xl/vbaProject.bin", FAKE_VBA
    )
    unprotect_workbook(macro_book, tmp_path / "out.xlsm")
    with zipfile.ZipFile(tmp_path / "out.xlsm") as after:
        assert after.read("xl/vbaProject.bin") == FAKE_VBA


def test_unhide_reveals_sheets_columns_and_reports_them(tmp_path: Path) -> None:
    """AC-18."""
    book = simple_workbook(
        tmp_path / "h.xlsx", hidden_sheet=True, hidden_columns=("B", "D"), protect=True
    )
    report = unprotect_workbook(book, tmp_path / "out.xlsx", unhide=True)
    workbook = openpyxl.load_workbook(tmp_path / "out.xlsx")

    assert all(sheet.sheet_state == "visible" for sheet in workbook.worksheets)
    demo = workbook["Demonstration"]
    assert not demo.column_dimensions["B"].hidden
    assert not demo.column_dimensions["D"].hidden
    assert report.unhidden_sheets == ["Internal", "Lookups"]
    assert report.unhidden_columns.get("Demonstration", 0) >= 1


def test_without_unhide_hidden_sheets_stay_hidden(tmp_path: Path) -> None:
    book = simple_workbook(tmp_path / "h.xlsx", hidden_sheet=True, protect=True)
    unprotect_workbook(book, tmp_path / "out.xlsx")
    workbook = openpyxl.load_workbook(tmp_path / "out.xlsx")
    assert workbook["Lookups"].sheet_state == "hidden"


def test_input_is_never_modified(tmp_path: Path) -> None:
    """AC-19."""
    book = simple_workbook(tmp_path / "src.xlsx", protect=True, hidden_columns=("C",))
    before = file_sha256(book)
    report = unprotect_workbook(book, tmp_path / "out.xlsx", unhide=True)
    assert file_sha256(book) == before
    assert report.source["sha256"] == before
    assert report.output["sha256"] != before


def test_default_output_path_is_derived(tmp_path: Path) -> None:
    book = simple_workbook(tmp_path / "derive.xlsx", protect=True)
    unprotect_workbook(book)
    assert (tmp_path / "derive.unprotected.xlsx").exists()


def test_refuses_to_overwrite_its_input(tmp_path: Path) -> None:
    book = simple_workbook(tmp_path / "same.xlsx", protect=True)
    with pytest.raises(UnsupportedWorkbookError) as excinfo:
        unprotect_workbook(book, book)
    assert "refusing" in str(excinfo.value)


def test_non_zip_input_is_rejected(tmp_path: Path) -> None:
    bogus = tmp_path / "bogus.xlsx"
    bogus.write_text("nope")
    with pytest.raises(UnsupportedWorkbookError):
        unprotect_workbook(bogus, tmp_path / "out.xlsx")


def test_unprotected_workbook_is_a_no_op_but_still_copies(tmp_path: Path) -> None:
    book = simple_workbook(tmp_path / "plain.xlsx", protect=False)
    report = unprotect_workbook(book, tmp_path / "out.xlsx")
    assert report.removed_sheet_protection == []
    assert (tmp_path / "out.xlsx").exists()
    assert openpyxl.load_workbook(tmp_path / "out.xlsx")["Demonstration"]["A4"].value


def test_report_summary_is_printable(tmp_path: Path) -> None:
    book = simple_workbook(tmp_path / "s.xlsx", protect=True)
    summary = unprotect_workbook(book, tmp_path / "out.xlsx").summary()
    assert "Demonstration" in summary
    assert "sheets unprotected: 1" in summary
