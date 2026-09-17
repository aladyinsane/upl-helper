"""Synthetic workbooks for testing the profiler.

These are not UPL templates. They exist to exercise specific structural
features — hidden sheets, protection, banners, multi-row headers — one at a
time. The realistic UPL fixture generator is a later spec.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

CANARY = "ZZCANARYZZ"
CANARY_NUMBER = 987654321.12

HEADERS = [
    "Provider Name",
    "Medicare Provider Number (CCN)",
    "Medicaid Days",
    "Cost Per Day",
    "Total Cost",
]


def simple_workbook(
    path: Path,
    *,
    banner: bool = True,
    rows: int = 10,
    hidden_sheet: bool = False,
    hidden_columns: tuple[str, ...] = (),
    protect: bool = False,
    protect_password: str | None = None,
    data_validation: bool = False,
    canary_data: bool = False,
    constant_in_formula_column: bool = False,
) -> Path:
    """Build a one-sheet workbook with an optional title banner above headers.

    Column E carries ``=C*D`` on every data row, which is what the formula
    pattern tests work against.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Demonstration"

    header_row = 1
    if banner:
        sheet["A1"] = "STATE OF EXAMPLE - INPATIENT HOSPITAL UPL DEMONSTRATION"
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
        sheet["A2"] = "SFY 2024"
        header_row = 4

    for offset, title in enumerate(HEADERS):
        sheet.cell(row=header_row, column=offset + 1, value=title)

    first_data_row = header_row + 1
    for index in range(rows):
        row = first_data_row + index
        sheet.cell(
            row=row,
            column=1,
            value=CANARY if canary_data else f"Example Hospital {index:02d}",
        )
        sheet.cell(row=row, column=2, value=CANARY if canary_data else f"14{index:04d}")
        sheet.cell(
            row=row, column=3, value=CANARY_NUMBER if canary_data else 1000 + index
        )
        sheet.cell(
            row=row, column=4, value=CANARY_NUMBER if canary_data else 850.25 + index
        )
        is_last = index == rows - 1
        if constant_in_formula_column and is_last:
            sheet.cell(row=row, column=5, value=123456.0)
        else:
            sheet.cell(row=row, column=5, value=f"=C{row}*D{row}")
        sheet.cell(row=row, column=5).number_format = "#,##0.00"

    for letter in hidden_columns:
        sheet.column_dimensions[letter].hidden = True

    if protect:
        sheet.protection.sheet = True
        if protect_password:
            sheet.protection.password = protect_password

    if data_validation:
        validation = DataValidation(
            type="list", formula1='"State,Non-State,Private"', allow_blank=True
        )
        sheet.add_data_validation(validation)
        validation.add(f"F{first_data_row}:F{first_data_row + rows}")

    if hidden_sheet:
        hidden = workbook.create_sheet("Lookups")
        hidden["A1"] = "Ownership"
        hidden["A2"] = "State"
        hidden.sheet_state = "hidden"
        very_hidden = workbook.create_sheet("Internal")
        very_hidden["A1"] = "Do not edit"
        very_hidden.sheet_state = "veryHidden"

    workbook.save(path)
    return path


def multi_row_header_workbook(path: Path, rows: int = 8) -> Path:
    """A two-row header: a category row above a field row.

    Both rows are mostly distinct text with numbers beneath, so both score
    highly and the adjacent-row rule joins them.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Demonstration"

    top = ["Provider", "Provider", "Utilization", "Rate", "Amount"]
    bottom = ["Name", "CCN", "Days", "Per Diem", "Total"]
    for offset, (upper, lower) in enumerate(zip(top, bottom, strict=True)):
        sheet.cell(row=1, column=offset + 1, value=upper)
        sheet.cell(row=2, column=offset + 1, value=lower)

    for index in range(rows):
        row = 3 + index
        sheet.cell(row=row, column=1, value=f"Hospital {index}")
        sheet.cell(row=row, column=2, value=f"14{index:04d}")
        sheet.cell(row=row, column=3, value=500 + index)
        sheet.cell(row=row, column=4, value=700.0 + index)
        sheet.cell(row=row, column=5, value=f"=C{row}*D{row}")

    workbook.save(path)
    return path


def empty_below_header_workbook(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Demonstration"
    for offset, title in enumerate(HEADERS):
        sheet.cell(row=1, column=offset + 1, value=title)
    workbook.save(path)
    return path
