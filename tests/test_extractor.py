"""Extraction (SPEC-0003 AC-1..4, AC-6..20)."""

from __future__ import annotations

import copy
from pathlib import Path

import pandas as pd
import pytest

from tests.builders import UPL_HEADERS, UPL_ROWS, upl_workbook
from tests.test_mapping_loader import MINIMAL
from upl_helper.extract.extractor import extract_path
from upl_helper.mapping.loader import load_mapping, parse_mapping
from upl_helper.mapping.models import MappingError
from upl_helper.model import FIELD_NAMES, SOURCE_SHEET_COLUMN

PROVISIONAL = "config/templates/inpatient-hospital-provisional.yaml"


def mapping():
    return load_mapping(PROVISIONAL)


def test_fields_resolve_from_aliases(tmp_path: Path) -> None:
    """AC-1."""
    book = upl_workbook(tmp_path / "wb.xlsx")
    frame = extract_path(book, mapping()).frame
    assert list(frame["provider_name"]) == [r[0] for r in UPL_ROWS]
    assert list(frame["ccn"]) == ["014001", "014002", "014003"]


def test_columns_resolve_by_header_not_position(tmp_path: Path) -> None:
    """AC-2. Same data, columns reversed, identical result."""
    straight = extract_path(upl_workbook(tmp_path / "a.xlsx"), mapping()).frame

    reversed_headers = list(reversed(UPL_HEADERS))
    reversed_rows = [tuple(reversed(row)) for row in UPL_ROWS]
    shuffled = extract_path(
        upl_workbook(tmp_path / "b.xlsx", headers=reversed_headers, rows=reversed_rows),
        mapping(),
    ).frame

    # assert_frame_equal rather than a list comparison: nulls are equal here,
    # and `nan != nan` would otherwise fail on the all-null optional columns.
    pd.testing.assert_frame_equal(straight, shuffled)


def test_missing_required_field_is_an_error(tmp_path: Path) -> None:
    """AC-3."""
    headers = [h for h in UPL_HEADERS if h != "UPL"]
    rows = [tuple(v for i, v in enumerate(r) if i != 5) for r in UPL_ROWS]
    book = upl_workbook(tmp_path / "wb.xlsx", headers=headers, rows=rows)
    with pytest.raises(MappingError) as excinfo:
        extract_path(book, mapping())
    assert "upl_amount" in str(excinfo.value)


def test_ambiguous_column_is_an_error_not_a_guess(tmp_path: Path) -> None:
    """AC-4."""
    headers = [*UPL_HEADERS, "Upper Payment Limit"]
    rows = [(*r, 1.0) for r in UPL_ROWS]
    book = upl_workbook(tmp_path / "wb.xlsx", headers=headers, rows=rows)
    with pytest.raises(MappingError) as excinfo:
        extract_path(book, mapping())
    message = str(excinfo.value)
    assert "more than one" in message
    assert "will not guess" in message


def test_unmapped_columns_are_reported(tmp_path: Path) -> None:
    """AC-6. A revised template that adds a column becomes visible."""
    book = upl_workbook(tmp_path / "wb.xlsx", extra_column="Brand New CMS Column")
    report = extract_path(book, mapping()).report
    unmapped = report.unmapped_columns["Cost Based Demonstration"]
    assert "Brand New CMS Column" in unmapped


def test_unmapped_sheets_are_reported(tmp_path: Path) -> None:
    """AC-7."""
    book = upl_workbook(tmp_path / "wb.xlsx", extra_sheet="Reviewer Notes")
    report = extract_path(book, mapping()).report
    assert "Reviewer Notes" in report.sheets_unmapped


def test_required_sheet_rule_matching_nothing_is_an_error(tmp_path: Path) -> None:
    """AC-8."""
    book = upl_workbook(tmp_path / "wb.xlsx", sheet_name="Something Else Entirely")
    with pytest.raises(MappingError) as excinfo:
        extract_path(book, mapping())
    assert "required but matched no sheet" in str(excinfo.value)


def test_stop_on_blank_key_halts_extraction(tmp_path: Path) -> None:
    """AC-9."""
    payload = copy.deepcopy(MINIMAL)
    rule = payload["sheets"][0]
    rule["match"] = {"name_matches": "Demonstration"}
    rule["data_rows"] = {"stop_on_blank_key": True}
    rule["fields"] = {
        "provider_name": {"aliases": ["provider name"]},
        "ccn": {"aliases": ["medicare provider number (ccn)"]},
        "ownership_category": {"aliases": ["ownership"]},
        "upl_amount": {"aliases": ["upl"]},
        "medicaid_payments_total": {"aliases": ["total medicaid payments"]},
    }
    rule["ownership"] = {
        "from": "column",
        "column": "ownership_category",
        "values": {
            "state": ["state"],
            "private": ["private"],
            "non_state_government": ["county"],
        },
    }

    book = upl_workbook(tmp_path / "wb.xlsx", blank_row_then_more=True)
    frame = extract_path(book, parse_mapping(payload)).frame
    assert "After The Gap" not in list(frame["provider_name"])
    assert len(frame) == 3


def test_continuing_past_a_blank_key_picks_up_later_rows(tmp_path: Path) -> None:
    book = upl_workbook(tmp_path / "wb.xlsx", blank_row_then_more=True)
    frame = extract_path(book, mapping()).frame
    assert "After The Gap" in list(frame["provider_name"])


def test_total_rows_are_excluded_and_counted(tmp_path: Path) -> None:
    """AC-10."""
    book = upl_workbook(tmp_path / "wb.xlsx", total_row=True)
    extraction = extract_path(book, mapping())
    assert "Total" not in list(extraction.frame["provider_name"])
    skipped = extraction.report.rows_skipped["Cost Based Demonstration"]
    assert skipped.get("total_row") == 1


def test_currency_and_percent_values_coerce(tmp_path: Path) -> None:
    """AC-11."""
    book = upl_workbook(tmp_path / "wb.xlsx")
    frame = extract_path(book, mapping()).frame
    assert frame.loc[0, "upl_amount"] == pytest.approx(45_000_000.0)
    assert frame.loc[0, "cost_to_charge_ratio"] == pytest.approx(0.412)


def test_unparseable_value_nulls_and_reports_with_a_cell_reference(
    tmp_path: Path,
) -> None:
    """AC-12. The point: it is reported, not silently dropped."""
    book = upl_workbook(tmp_path / "wb.xlsx", bad_value_cell=(2, 6, "see footnote"))
    extraction = extract_path(book, mapping())
    assert (
        extraction.frame.loc[0, "upl_amount"] != extraction.frame.loc[0, "upl_amount"]
    )

    issues = [i for i in extraction.report.coercion_issues if i.field == "upl_amount"]
    assert len(issues) == 1
    assert issues[0].cell == "F2"
    assert issues[0].raw == "see footnote"
    assert issues[0].sheet == "Cost Based Demonstration"


def test_leading_zero_in_ccn_survives(tmp_path: Path) -> None:
    """AC-13."""
    frame = extract_path(upl_workbook(tmp_path / "wb.xlsx"), mapping()).frame
    assert frame.loc[0, "ccn"] == "014001"


def test_ownership_categories_map_to_canonical_values(tmp_path: Path) -> None:
    """AC-14."""
    frame = extract_path(upl_workbook(tmp_path / "wb.xlsx"), mapping()).frame
    assert list(frame["ownership_category"]) == [
        "state",
        "non_state_government",
        "private",
    ]


def test_unrecognized_ownership_becomes_unknown_and_is_reported(
    tmp_path: Path,
) -> None:
    """AC-14, the failure half. Ownership decides the UPL group; never guess."""
    rows = [("Odd One", "014009", "Martian Collective", 1, "0.5", 1.0, 1.0)]
    book = upl_workbook(tmp_path / "wb.xlsx", rows=rows)
    extraction = extract_path(book, mapping())
    assert extraction.frame.loc[0, "ownership_category"] == "unknown"
    assert any(
        i.field == "ownership_category" for i in extraction.report.coercion_issues
    )


def test_absent_optional_fields_still_have_columns(tmp_path: Path) -> None:
    """AC-16. Checks never have to ask whether a column exists."""
    extraction = extract_path(upl_workbook(tmp_path / "wb.xlsx"), mapping())
    for name in FIELD_NAMES:
        assert name in extraction.frame.columns
    assert extraction.frame["npi"].isna().all()
    assert "npi" in extraction.report.fields_missing["Cost Based Demonstration"]


def test_provenance_matches_the_frame_and_points_at_cells(tmp_path: Path) -> None:
    """AC-17."""
    extraction = extract_path(upl_workbook(tmp_path / "wb.xlsx"), mapping())
    assert len(extraction.provenance) == len(extraction.frame)
    assert list(extraction.provenance.columns) == list(FIELD_NAMES)
    assert extraction.provenance.loc[0, "ccn"] == "Cost Based Demonstration!B2"
    assert extraction.provenance.loc[2, "upl_amount"] == "Cost Based Demonstration!F4"


def test_extraction_is_repeatable(tmp_path: Path) -> None:
    """AC-18."""
    book = upl_workbook(tmp_path / "wb.xlsx")
    first = extract_path(book, mapping()).frame
    second = extract_path(book, mapping()).frame
    assert first.equals(second)


def test_source_sheet_is_preserved(tmp_path: Path) -> None:
    """AC-19."""
    extraction = extract_path(upl_workbook(tmp_path / "wb.xlsx"), mapping())
    assert set(extraction.frame[SOURCE_SHEET_COLUMN]) == {"Cost Based Demonstration"}


def test_upl_gap_is_not_computed_by_the_extractor(tmp_path: Path) -> None:
    """AC-20. Recomputing it here would make ARI004 tautological."""
    extraction = extract_path(upl_workbook(tmp_path / "wb.xlsx"), mapping())
    assert extraction.frame["upl_gap"].isna().all()


def test_extraction_reads_calculated_values_not_formula_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Some real templates compute every demonstration cell by formula from a
    separate input sheet (the inpatient hospital methodology tabs do). If
    extraction opened the workbook the way the profiler does -- formulas as
    text -- every field on such a template would extract as a formula
    string instead of a value. Pin the wiring: extraction must open with
    data_only=True.
    """
    import upl_helper.profile.loader as loader_module

    original_load_workbook = loader_module.openpyxl.load_workbook
    seen: dict[str, object] = {}

    def spy(*args: object, **kwargs: object):
        seen.update(kwargs)
        return original_load_workbook(*args, **kwargs)

    monkeypatch.setattr(loader_module.openpyxl, "load_workbook", spy)
    extract_path(upl_workbook(tmp_path / "wb.xlsx"), mapping())
    assert seen["data_only"] is True


def test_report_summary_is_printable(tmp_path: Path) -> None:
    book = upl_workbook(tmp_path / "wb.xlsx", extra_sheet="Notes", total_row=True)
    summary = extract_path(book, mapping()).report.summary()
    assert "extracted 3 row(s)" in summary
    assert "Notes" in summary


def _workbook_with_instructions_gap(path: Path) -> Path:
    """Header at row 1, a blank row 2, an instructions row at row 3, data at row 4.

    Mirrors the real inpatient hospital UPL template's methodology tabs: a
    blank row and an instructions row separate the header from the first
    real provider row, rather than data starting immediately below it.
    """
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Demo"
    sheet["A1"] = "provider name"
    sheet["B1"] = "ccn"
    sheet["A3"] = "Obtain from Medicare Cost Reports"
    sheet["B3"] = "Obtain from Medicare Cost Reports"
    sheet["A4"] = "Example Hospital"
    sheet["B4"] = "140001"
    workbook.save(path)
    return path


def _gap_mapping(start_row: int | None) -> dict:
    payload = copy.deepcopy(MINIMAL)
    payload["sheets"][0]["fields"] = {
        "provider_name": {"aliases": ["provider name"]},
        "ccn": {"aliases": ["ccn"]},
    }
    payload["sheets"][0]["key_field"] = "ccn"
    payload["sheets"][0]["header"] = {"rows": [1]}
    if start_row is not None:
        payload["sheets"][0]["data_rows"] = {"start_row": start_row}
    return payload


def test_start_row_override_skips_a_gap_between_header_and_data(
    tmp_path: Path,
) -> None:
    book = _workbook_with_instructions_gap(tmp_path / "wb.xlsx")
    frame = extract_path(book, parse_mapping(_gap_mapping(start_row=4))).frame
    assert list(frame["provider_name"]) == ["Example Hospital"]


def test_without_start_row_a_gap_stops_extraction_before_any_real_data(
    tmp_path: Path,
) -> None:
    """The failure mode `start_row` exists to fix, pinned down as a test.

    Without the override, extraction assumes data starts immediately below
    the header, hits the blank row 2 first, and (with the default
    stop_on_blank_key) stops there -- silently extracting zero providers
    from a workbook that has one.
    """
    book = _workbook_with_instructions_gap(tmp_path / "wb.xlsx")
    frame = extract_path(book, parse_mapping(_gap_mapping(start_row=None))).frame
    assert frame.empty
