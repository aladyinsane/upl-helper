"""Per-column structural profiling.

Columns are profiled in a single pass over the sheet rather than one pass per
column, because a real template is wide and re-scanning it per column is
quadratic. The public entry point is therefore plural.
"""

from __future__ import annotations

import datetime as dt
from collections import Counter
from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import get_column_letter

from upl_helper.profile.models import CellTypeCounts, ColumnProfile, FormulaPattern
from upl_helper.profile.text import normalize_header, templatize_formula

HEADER_JOIN = " / "


@dataclass
class _ColumnAccumulator:
    index: int
    header_parts: list[str] = field(default_factory=list)
    counts: CellTypeCounts = field(default_factory=CellTypeCounts)
    formula_counts: Counter[str] = field(default_factory=Counter)
    formula_first_row: dict[str, int] = field(default_factory=dict)
    formula_last_row: dict[str, int] = field(default_factory=dict)
    number_formats: Counter[str] = field(default_factory=Counter)


def _classify(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, str):
        if value.startswith("="):
            return "formula"
        if value.strip() == "":
            return "blank"
        if value.startswith("#") and value.endswith(("!", "?", "A", "0")):
            return "error"
        return "string"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return "date"
    if isinstance(value, (int, float)):
        return "numeric"
    return "string"


def profile_columns(worksheet: Any, header_rows: list[int]) -> list[ColumnProfile]:
    """Profile every column in the sheet's used range."""
    min_col = worksheet.min_column or 1
    max_col = worksheet.max_column or 1
    min_row = worksheet.min_row or 1
    max_row = worksheet.max_row or 1

    accumulators = {
        col: _ColumnAccumulator(index=col) for col in range(min_col, max_col + 1)
    }
    header_row_set = set(header_rows)
    last_header_row = max(header_rows) if header_rows else min_row - 1

    for row in worksheet.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            acc = accumulators.get(cell.column)
            if acc is None:
                continue
            value = cell.value

            if cell.row in header_row_set:
                if value is not None and str(value).strip() != "":
                    acc.header_parts.append(str(value).strip())
                continue

            if cell.row <= last_header_row:
                # Title banners and notes above the header describe nothing.
                continue

            kind = _classify(value)
            setattr(acc.counts, kind, getattr(acc.counts, kind) + 1)

            if kind == "formula":
                pattern = templatize_formula(str(value), cell.row)
                acc.formula_counts[pattern] += 1
                acc.formula_first_row.setdefault(pattern, cell.row)
                acc.formula_last_row[pattern] = cell.row

            if kind != "blank" and cell.number_format:
                acc.number_formats[cell.number_format] += 1

    profiles: list[ColumnProfile] = []
    for col in range(min_col, max_col + 1):
        acc = accumulators[col]
        letter = get_column_letter(col)
        dimension = worksheet.column_dimensions.get(letter)

        header_text = HEADER_JOIN.join(acc.header_parts) if acc.header_parts else None
        patterns = [
            FormulaPattern(
                pattern=pattern,
                count=count,
                first_row=acc.formula_first_row[pattern],
                last_row=acc.formula_last_row[pattern],
            )
            for pattern, count in sorted(
                acc.formula_counts.items(), key=lambda kv: (-kv[1], kv[0])
            )
        ]
        number_format = (
            acc.number_formats.most_common(1)[0][0] if acc.number_formats else None
        )

        profiles.append(
            ColumnProfile(
                letter=letter,
                index=col,
                hidden=bool(dimension.hidden) if dimension is not None else False,
                width=float(dimension.width)
                if dimension is not None and dimension.width is not None
                else None,
                header_text=header_text,
                header_normalized=normalize_header(header_text),
                number_format=number_format,
                cell_type_counts=acc.counts,
                formula_patterns=patterns,
            )
        )
    return profiles
