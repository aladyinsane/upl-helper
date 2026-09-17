"""Opening a workbook and recording what file it came from."""

from __future__ import annotations

import hashlib
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from upl_helper.errors import UnsupportedWorkbookError

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}


@dataclass
class LoadedWorkbook:
    workbook: Any
    path: Path
    sha256: str
    size_bytes: int


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_workbook_structure(path: str | Path) -> LoadedWorkbook:
    """Open a workbook for structural inspection.

    Formulas are read as text rather than as cached values, because the
    formulas are what the profiler is interested in. ``keep_vba`` preserves
    macro-enabled workbooks.
    """
    path = Path(path)

    if not path.exists():
        raise UnsupportedWorkbookError(f"{path}: no such file")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise UnsupportedWorkbookError(
            f"{path}: unsupported extension {path.suffix!r}; "
            f"expected one of {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )
    if not zipfile.is_zipfile(path):
        raise UnsupportedWorkbookError(
            f"{path}: not a readable .xlsx/.xlsm file (it is not a zip archive). "
            "Legacy .xls workbooks are not supported; re-save as .xlsx first."
        )

    try:
        workbook = openpyxl.load_workbook(
            path, data_only=False, keep_vba=path.suffix.lower() in {".xlsm", ".xltm"}
        )
    except Exception as exc:  # openpyxl raises a wide variety here
        raise UnsupportedWorkbookError(f"{path}: could not be opened ({exc})") from exc

    return LoadedWorkbook(
        workbook=workbook,
        path=path,
        sha256=file_sha256(path),
        size_bytes=path.stat().st_size,
    )
