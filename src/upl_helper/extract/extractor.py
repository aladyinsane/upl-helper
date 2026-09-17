"""Workbook plus mapping -> canonical frame."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.utils import get_column_letter

from upl_helper.extract.coerce import coerce_ownership, coerce_value
from upl_helper.mapping.models import SheetRule, TemplateMapping
from upl_helper.mapping.resolve import ResolvedSheet, resolve_sheets
from upl_helper.model import (
    FIELD_NAMES,
    FIELDS_BY_NAME,
    SOURCE_SHEET_COLUMN,
    Dtype,
    ExtractionContext,
    Ownership,
)
from upl_helper.profile.loader import load_workbook_structure
from upl_helper.profile.text import normalize_header

_PANDAS_DTYPES: dict[Dtype, str] = {
    Dtype.STRING: "object",
    Dtype.INTEGER: "Int64",
    Dtype.NUMBER: "float64",
    Dtype.DATE: "object",
    Dtype.RATIO: "float64",
    Dtype.MONEY: "float64",
}


@dataclass(frozen=True)
class ReportedIssue:
    sheet: str
    cell: str
    field: str
    dtype: str
    raw: str
    message: str


@dataclass
class ExtractionReport:
    sheets_matched: list[dict[str, Any]] = field(default_factory=list)
    sheets_unmapped: list[str] = field(default_factory=list)
    rules_unmatched: list[str] = field(default_factory=list)
    fields_missing: dict[str, list[str]] = field(default_factory=dict)
    unmapped_columns: dict[str, list[str]] = field(default_factory=dict)
    rows_skipped: dict[str, dict[str, int]] = field(default_factory=dict)
    coercion_issues: list[ReportedIssue] = field(default_factory=list)

    def summary(self) -> str:
        rows = sum(s["rows_extracted"] for s in self.sheets_matched)
        lines = [f"extracted {rows} row(s) from {len(self.sheets_matched)} sheet(s)"]
        for matched in self.sheets_matched:
            lines.append(f"  {matched['sheet']}: {matched['rows_extracted']} row(s)")
        if self.sheets_unmapped:
            lines.append(
                f"sheets not covered by any rule: {', '.join(self.sheets_unmapped)}"
            )
        for sheet, columns in self.unmapped_columns.items():
            if columns:
                lines.append(f"unmapped columns on {sheet}: {', '.join(columns)}")
        for sheet, missing in self.fields_missing.items():
            if missing:
                lines.append(f"fields absent on {sheet}: {', '.join(missing)}")
        if self.coercion_issues:
            lines.append(f"coercion issues: {len(self.coercion_issues)}")
            for issue in self.coercion_issues[:10]:
                lines.append(
                    f"  {issue.sheet}!{issue.cell} {issue.field}: "
                    f"{issue.message} ({issue.raw!r})"
                )
            if len(self.coercion_issues) > 10:
                lines.append(f"  ... and {len(self.coercion_issues) - 10} more")
        return "\n".join(lines)


@dataclass
class Extraction:
    frame: pd.DataFrame
    provenance: pd.DataFrame
    report: ExtractionReport
    context: ExtractionContext | None = None


def _empty_frame() -> pd.DataFrame:
    frame = pd.DataFrame({name: pd.Series(dtype="object") for name in FIELD_NAMES})
    frame[SOURCE_SHEET_COLUMN] = pd.Series(dtype="object")
    return frame


def _is_total_row(values: dict[str, object], rule: SheetRule) -> bool:
    excluded = {normalize_header(v) for v in rule.data_rows.exclude_if_normalized_in}
    for marker in rule.data_rows.total_row_markers:
        raw = values.get(marker)
        if raw is None:
            continue
        if normalize_header(str(raw)) in excluded:
            return True
    return False


def _extract_sheet(
    resolved: ResolvedSheet, report: ExtractionReport
) -> tuple[list[dict[str, object]], list[dict[str, str]]]:
    worksheet = resolved.worksheet
    rule = resolved.rule
    columns = resolved.resolution.columns
    sheet_name = resolved.name

    first_row = (max(resolved.header_rows) if resolved.header_rows else 0) + 1
    last_row = worksheet.max_row or 0
    skipped = defaultdict(int)

    records: list[dict[str, object]] = []
    provenance_rows: list[dict[str, str]] = []

    for row_index in range(first_row, last_row + 1):
        raw_values = {
            name: worksheet.cell(row=row_index, column=col).value
            for name, col in columns.items()
        }

        if all(v is None or str(v).strip() == "" for v in raw_values.values()):
            skipped["blank"] += 1
            # A wholly blank row has a blank key too. When the sheet is
            # configured to stop there, stop: what follows a gap in a CMS
            # template is usually footnotes, not more providers.
            if rule.data_rows.stop_on_blank_key:
                break
            continue

        if _is_total_row(raw_values, rule):
            skipped["total_row"] += 1
            continue

        key_raw = raw_values.get(rule.key_field)
        if key_raw is None or str(key_raw).strip() == "":
            skipped["blank_key"] += 1
            if rule.data_rows.stop_on_blank_key:
                break
            continue

        record: dict[str, object] = {}
        provenance: dict[str, str] = {}

        for name, col in columns.items():
            cell_ref = f"{get_column_letter(col)}{row_index}"
            raw = raw_values[name]
            dtype = rule.fields[name].dtype

            if name == "ownership_category" and rule.ownership is not None:
                value, issue = coerce_ownership(raw, rule.ownership.values)
            else:
                value, issue = coerce_value(raw, dtype, name)

            record[name] = value
            if raw is not None and str(raw).strip() != "":
                provenance[name] = f"{sheet_name}!{cell_ref}"
            if issue is not None:
                report.coercion_issues.append(
                    ReportedIssue(
                        sheet=sheet_name,
                        cell=cell_ref,
                        field=issue.field or name,
                        dtype=str(issue.dtype),
                        raw=issue.raw,
                        message=issue.message,
                    )
                )

        if rule.ownership is not None and rule.ownership.source == "constant":
            record["ownership_category"] = rule.ownership.constant
        record.setdefault("ownership_category", Ownership.UNKNOWN.value)

        record[SOURCE_SHEET_COLUMN] = sheet_name
        records.append(record)
        provenance_rows.append(provenance)

    report.rows_skipped[sheet_name] = dict(skipped)
    return records, provenance_rows


def extract(
    workbook: Any, mapping: TemplateMapping, context: ExtractionContext | None = None
) -> Extraction:
    resolved_sheets, unmapped, unmatched_rules = resolve_sheets(workbook, mapping)
    report = ExtractionReport(sheets_unmapped=unmapped, rules_unmatched=unmatched_rules)

    all_records: list[dict[str, object]] = []
    all_provenance: list[dict[str, str]] = []

    for resolved in resolved_sheets:
        if resolved.rule.role != "demonstration":
            continue
        records, provenance = _extract_sheet(resolved, report)
        all_records.extend(records)
        all_provenance.extend(provenance)
        report.sheets_matched.append(
            {
                "rule": resolved.rule.label,
                "sheet": resolved.name,
                "rows_extracted": len(records),
            }
        )
        report.fields_missing[resolved.name] = resolved.resolution.missing_fields
        report.unmapped_columns[resolved.name] = resolved.resolution.unmapped_columns

    frame = _empty_frame() if not all_records else pd.DataFrame(all_records)

    # Optional fields absent from the workbook still get a column, so no check
    # ever has to ask whether a column exists.
    for name in FIELD_NAMES:
        if name not in frame.columns:
            frame[name] = pd.Series([None] * len(frame), dtype="object")
    frame = frame[[*FIELD_NAMES, SOURCE_SHEET_COLUMN]]

    for name in FIELD_NAMES:
        target = _PANDAS_DTYPES[FIELDS_BY_NAME[name].dtype]
        if target in {"float64", "Int64"}:
            frame[name] = pd.to_numeric(frame[name], errors="coerce").astype(target)

    provenance = pd.DataFrame(all_provenance, columns=list(FIELD_NAMES))
    provenance = provenance.reindex(columns=list(FIELD_NAMES))
    frame = frame.reset_index(drop=True)
    provenance = provenance.reset_index(drop=True)

    return Extraction(
        frame=frame, provenance=provenance, report=report, context=context
    )


def extract_path(
    workbook_path: str | Path,
    mapping: TemplateMapping,
    context: ExtractionContext | None = None,
) -> Extraction:
    loaded = load_workbook_structure(workbook_path)
    return extract(loaded.workbook, mapping, context)
